"""多格式文档加载：把 data/ 下的文件统一转成 LangChain Document。

解析策略（2026-08 优化）：
- PDF 优先走 PyMuPDF4LLM 布局感知解析：保留标题层级、段落、表格（Markdown），
  并带页码；失败时回退 PyPDFLoader（纯文本按页）。
- 扫描件/图片页（几乎无文本）可选用视觉模型（SenseNova）做 OCR，默认关闭。
- EPUB 按章节分组；.doc 用 doc2txt/legacy-doc；Excel 按 sheet 逐行管道连接。
"""

from __future__ import annotations

import re
from pathlib import Path

from langchain_community.document_loaders import (
    CSVLoader,
    Docx2txtLoader,
    PyPDFLoader,
    TextLoader,
)
from langchain_core.documents import Document

TEXT_SUFFIXES = {".txt", ".md", ".markdown"}
SUPPORTED_SUFFIXES = TEXT_SUFFIXES | {".csv", ".doc", ".docx", ".xlsx", ".pdf", ".epub"}


def list_data_files(data_dir: Path) -> list[Path]:
    """递归列出 data/ 下所有受支持格式的文件。"""
    return sorted(
        p
        for p in Path(data_dir).rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES
    )


def load_xlsx(file_path: Path) -> list[Document]:
    """Excel 专用加载器：每个 sheet 生成一个 Document，每行用 | 连接单元格。"""
    from openpyxl import load_workbook  # 延迟导入

    workbook = load_workbook(file_path, data_only=True, read_only=True)
    docs: list[Document] = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [
                    str(cell).strip()
                    for cell in row
                    if cell is not None and str(cell).strip()
                ]
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                docs.append(
                    Document(
                        page_content="\n".join(rows),
                        metadata={"source": file_path.name, "sheet": sheet.title},
                    )
                )
    finally:
        workbook.close()
    return docs


def load_documents(
    data_dir: Path,
    rel_paths: set[str] | None = None,
    *,
    layout_pdf: bool = True,
    pdf_workers: int = 0,
    vision=None,
    vision_ocr_pages: int = 0,
) -> list[Document]:
    """加载文档，按扩展名分发解析器。

    rel_paths：只加载指定相对路径的文件（增量索引时只读新文件）；
    None 时加载全部文件（全量重建时用）。
    layout_pdf：是否使用布局感知 PDF 解析（PyMuPDF4LLM）。
    pdf_workers：大 PDF 布局解析并行进程数（0=串行）。
    vision：SenseNova 视觉客户端（可选，用于扫描页 OCR）。
    vision_ocr_pages：单个 PDF 最多 OCR 页数，0 表示关闭。
    """
    files = list_data_files(data_dir)
    if rel_paths is not None:
        files = [p for p in files if p.relative_to(data_dir).as_posix() in rel_paths]
    if not files:
        return []

    docs: list[Document] = []
    for file_path in files:
        suffix = file_path.suffix.lower()
        try:
            if suffix in TEXT_SUFFIXES:
                loaded = TextLoader(
                    str(file_path), encoding="utf-8", autodetect_encoding=True
                ).load()
            elif suffix == ".csv":
                loaded = CSVLoader(str(file_path), encoding="utf-8").load()
            elif suffix == ".docx":
                loaded = Docx2txtLoader(str(file_path)).load()
            elif suffix == ".pdf":
                loaded = load_pdf(
                    file_path,
                    layout=layout_pdf,
                    workers=pdf_workers,
                    vision=vision,
                    vision_ocr_pages=vision_ocr_pages,
                )
            elif suffix == ".epub":
                loaded = load_epub(file_path)
            elif suffix == ".doc":
                loaded = load_doc(file_path)
            elif suffix == ".xlsx":
                loaded = load_xlsx(file_path)
            else:
                continue
            # 统一补相对路径元数据：检索结果据此定位原文（子目录重名文档不混淆）
            rel_path = file_path.relative_to(data_dir).as_posix()
            for doc in loaded:
                doc.metadata.setdefault("relative_path", rel_path)
            docs.extend(loaded)
        except ImportError as exc:
            print(f"  跳过 {file_path.name}：缺少依赖 {exc.name}")
        except Exception as exc:
            print(f"  跳过 {file_path.name}：{exc}")
    return docs


def load_pdf(
    file_path: Path | str,
    *,
    layout: bool = True,
    workers: int = 0,
    vision=None,
    vision_ocr_pages: int = 0,
) -> list[Document]:
    """PDF 加载：布局感知解析 + 可选扫描页 OCR。

    PyMuPDF4LLM 会输出带标题层级与表格的 Markdown，页面间阅读顺序也更合理，
    显著改善后续按段落分组的 Parent-Child 切分质量。
    """
    file_path = Path(file_path)
    docs: list[Document] = []
    if layout:
        try:
            if workers > 0 and _pdf_page_count(str(file_path)) >= 300:
                docs = _layout_markdown_parallel(str(file_path), workers)
            else:
                docs = _layout_markdown_serial(str(file_path))
        except ImportError:
            docs = []
        except Exception as exc:
            print(f"  PyMuPDF4LLM 解析 {file_path.name} 失败，回退 PyPDFLoader：{exc}")
            docs = []

    if not docs:
        loaded = PyPDFLoader(str(file_path)).load()
        for doc in loaded:
            doc.metadata.setdefault("source", file_path.name)
        docs = loaded

    if vision is not None and vision_ocr_pages > 0:
        docs = _ocr_sparse_pdf_pages(file_path, docs, vision, vision_ocr_pages)
    return docs


def _pdf_page_count(file_path: str) -> int:
    """快速获取 PDF 页数（不加载版面模型）。"""
    import fitz

    with fitz.open(file_path) as doc:
        return doc.page_count


def _layout_markdown_serial(file_path: str) -> list[Document]:
    """串行布局解析：整本一次转 Markdown（带页码）。"""
    import pymupdf4llm

    pages = pymupdf4llm.to_markdown(
        file_path,
        page_chunks=True,
        show_progress=False,
        write_images=False,
    )
    docs: list[Document] = []
    name = Path(file_path).name
    for page in pages:
        text = (page.get("text") or "").strip()
        if not text:
            continue
        meta = dict(page.get("metadata") or {})
        docs.append(
            Document(
                page_content=text,
                metadata={"source": name, "page": meta.get("page_number") or 1},
            )
        )
    return docs


def _layout_markdown_parallel(file_path: str, workers: int) -> list[Document]:
    """按页范围多进程布局解析（布局推理是进程内全局状态，线程并行无效）。"""
    import math
    from concurrent.futures import ProcessPoolExecutor, as_completed

    total = _pdf_page_count(file_path)
    chunk = max(50, math.ceil(total / max(1, workers)))
    ranges = [
        (file_path, start, min(start + chunk, total))
        for start in range(0, total, chunk)
    ]
    docs: list[Document] = []
    failed = 0
    with ProcessPoolExecutor(max_workers=max(1, workers)) as pool:
        futures = {pool.submit(_layout_range, *r): r for r in ranges}
        for future in as_completed(futures):
            try:
                docs.extend(future.result())
            except Exception as exc:
                failed += 1
                print(f"  PDF 并行解析段失败：{exc}")
    if failed:
        print(f"  并行解析有 {failed} 段失败，回退串行解析整本")
        return _layout_markdown_serial(file_path)
    docs.sort(key=lambda d: d.metadata.get("page") or 0)
    return docs


def _layout_range(file_path: str, start: int, end: int) -> list[Document]:
    """子进程任务：解析 [start, end) 页范围（0 基页码）。"""
    import pymupdf4llm

    pages = pymupdf4llm.to_markdown(
        file_path,
        page_chunks=True,
        show_progress=False,
        write_images=False,
        pages=list(range(start, end)),
    )
    docs: list[Document] = []
    name = Path(file_path).name
    for page in pages:
        text = (page.get("text") or "").strip()
        if not text:
            continue
        meta = dict(page.get("metadata") or {})
        docs.append(
            Document(
                page_content=text,
                metadata={"source": name, "page": meta.get("page_number") or start + 1},
            )
        )
    return docs


def _ocr_sparse_pdf_pages(
    file_path: Path,
    docs: list[Document],
    vision,
    max_pages: int,
) -> list[Document]:
    """找出几乎无文本的页面（扫描件），用视觉模型 OCR 并追加为文档。"""
    try:
        import fitz  # pymupdf
    except ImportError:
        return docs

    empty_pages: list[int] = []
    for doc in docs:
        page = doc.metadata.get("page")
        text = (doc.page_content or "").strip()
        if page and len(text) < 30:
            empty_pages.append(int(page))
    if not empty_pages:
        return docs

    done = 0
    pdf = fitz.open(str(file_path))
    try:
        for page_no in sorted(set(empty_pages)):
            if done >= max_pages:
                break
            try:
                page = pdf.load_page(page_no - 1)
                pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
                raw = pix.tobytes("png")
                text = vision.describe_image_bytes(
                    raw,
                    filename=f"{file_path.name}_p{page_no}.png",
                    question=(
                        "这是 PDF 的第 %d 页扫描件。请完整识别页面内容，"
                        "包括所有标题、正文与表格，输出为易于阅读的纯文本。"
                        % page_no
                    ),
                )
                if text:
                    docs.append(
                        Document(
                            page_content=f"[扫描页 OCR] {text}",
                            metadata={
                                "source": file_path.name,
                                "page": page_no,
                                "ocr": True,
                            },
                        )
                    )
                    done += 1
            except Exception as exc:
                print(f"  OCR 第 {page_no} 页失败：{exc}")
    finally:
        pdf.close()
    if done:
        print(f"  已用视觉模型 OCR {file_path.name} 的 {done} 个扫描页")
    return docs


def load_doc(file_path: Path | str) -> list[Document]:
    """旧版二进制 .doc（Word 97-2003）加载器。

    优先使用 doc2txt（内置 antiword 二进制，跨平台、鲁棒），
    失败时退回 legacy-doc（纯 Python 实现，适合较小的文件）。
    """
    file_path = Path(file_path)
    text = ""
    errors: list[str] = []
    try:
        from doc2txt import extract_text as doc2txt_extract

        text = doc2txt_extract(str(file_path))
    except Exception as exc:  # noqa: BLE001
        errors.append(f"doc2txt: {exc}")
        try:
            from legacy_doc import extract_text as legacy_extract

            result = legacy_extract(file_path.read_bytes())
            text = result.text if hasattr(result, "text") else str(result)
        except Exception as exc2:  # noqa: BLE001
            errors.append(f"legacy-doc: {exc2}")
    if not text:
        raise ValueError(f"无法解析 .doc 文件（{'；'.join(errors) or '空文本'}）")

    # 轻量归一化：统一换行、去掉零宽字符、压缩多余空行
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\ufeff", "")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return [Document(page_content=text, metadata={"source": file_path.name})]


def load_epub(file_path: Path | str) -> list[Document]:
    """EPUB 加载器：提取全部文本内容，按章节分组。

    EPUB 本质是一个 ZIP 包，内含 XHTML 格式的章节文件。
    用 ebooklib 解包 + BeautifulSoup 提取纯文本。
    """
    file_path = Path(file_path)
    import ebooklib
    from bs4 import BeautifulSoup
    from ebooklib import epub

    book = epub.read_epub(str(file_path))
    docs: list[Document] = []
    for item in book.get_items():
        if item.get_type() != ebooklib.ITEM_DOCUMENT:
            continue
        # item.get_content() 返回 bytes，用 BeautifulSoup 提取纯文本
        soup = BeautifulSoup(item.get_content(), "html.parser")
        text = soup.get_text(separator="\n", strip=True)
        if not text:
            continue
        # 用文件名作为章节标识
        chapter_name = item.get_name() or "未知章节"
        docs.append(
            Document(
                page_content=text,
                metadata={"source": file_path.name, "chapter": chapter_name},
            )
        )
    if not docs:
        print(f"  警告：{file_path.name} 未提取到任何文本内容")
    else:
        print(f"  已加载 ePub：{file_path.name}，共 {len(docs)} 个章节")
    return docs
